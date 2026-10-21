import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE = Path(r"C:\Users\User\Downloads\Project-autofilter")
CONS = "C:/Users/User/Downloads/CONSIDERA\u00c7\u00d5ES/"


def norm(v):
    if pd.isna(v):
        return ""
    try:
        n = float(v)
        if n == int(n):
            return str(int(n))
    except (TypeError, ValueError):
        pass
    return str(v).strip()


est = pd.read_excel(CONS + "Estoque CMC.xlsx")
s2 = {}
for _, r in est.iterrows():
    c = norm(r["Cód Confec"])
    if c and pd.notna(r["Saldo 2"]):
        s2[c] = float(r["Saldo 2"])

wb = load_workbook(BASE / "_preview_excel_qe.xlsx")
ws = wb.worksheets[0]

erros = 0
total = 0
limitados = 0
for i in range(6, 33):
    cod = norm(ws.cell(row=i, column=2).value)
    qe = ws.cell(row=i, column=6).value
    saldo = ws.cell(row=i, column=5).value
    bruto = math.ceil(abs(s2[cod]) * 1.10)
    esperado = min(bruto, int(saldo))
    total += qe
    if esperado == int(saldo) and bruto > int(saldo):
        limitados += 1
    if esperado != qe:
        erros += 1
        print(f"XX cod={cod} excel={qe} esperado={esperado} (bruto={bruto}, saldo={saldo})")

print(f"linhas verificadas: 27 | erros: {erros}")
print(f"linhas limitadas pelo Saldo do pedido: {limitados} | sem limite (bruto < saldo): {27 - limitados}")
print(f"Q.E total na planilha: {ws.cell(row=33, column=6).value} | somado: {total}")
print(f"Saldo 2 do estoque usado: {[s2[norm(ws.cell(row=i, column=2).value)] for i in (6, 7, 8)]}")
